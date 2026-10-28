"""Build the progress deck (report/day1_report.pptx) from results/day1,
results/day2 (partition x workload grid, 16K profile rerun) and
results/repeat (3 repeats of the key sweep points).

Needs python-pptx, matplotlib, pandas, plus the vllm010 env for
hybrid_pp.analyze_profile. Re-run after new results land:

    python report/make_report.py --profile-python /venv/vllm010/bin/python
"""

import argparse
import glob
import json
import os
import subprocess
import sys
import tempfile

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import pandas as pd  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402
from pptx.util import Emu, Inches, Pt  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
from hybrid_pp.analyze_memory import parse_log, summarize  # noqa: E402

RES = os.path.join(ROOT, "results", "day1")
RES2 = os.path.join(ROOT, "results", "day2")
REP = os.path.join(ROOT, "results", "repeat")
GRID_PARTS = [("16-40", 16), ("20-36", 20), ("24-32", 24), ("default", 28),
              ("32-24", 32), ("36-20", 36), ("40-16", 40)]
GRID_WLS = [("1024-256", "短对话 1K→256"), ("512-1024", "长输出 512→1024"),
            ("8192-16", "长输入 8K→16"), ("16384-128", "长上下文 16K→128")]
C_SERIES = ["#2563EB", "#EA580C", "#0D9488", "#9333EA"]
PATTERN = "M-M-M-MM-M-M-M*-M-M-M*-M-M-M-M*-M-M-M-M*-M-MM-M-M-M-M-M-"

# palette
INK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x6B, 0x72, 0x80)
ACCENT = RGBColor(0x25, 0x63, 0xEB)
GOOD = RGBColor(0x05, 0x96, 0x69)
WARN = RGBColor(0xD9, 0x77, 0x06)
BAD = RGBColor(0xDC, 0x26, 0x26)
LIGHT = RGBColor(0xF3, 0xF4, 0xF6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
C_MAMBA, C_MLP, C_ATTN = "#0F766E", "#9CA3AF", "#EA580C"
C_W, C_ACT, C_KV, C_UNUSED = "#64748B", "#A78BFA", "#2563EB", "#F87171"

FONT_EA = "Microsoft YaHei"
FONT_LATIN = "Calibri"
# Noto on Linux, Microsoft YaHei on Windows; matplotlib uses the first found.
plt.rcParams["font.family"] = ["Noto Sans CJK SC", "Microsoft YaHei",
                               "DejaVu Sans"]
plt.rcParams["axes.spines.top"] = False
plt.rcParams["axes.spines.right"] = False

W, H = Inches(13.333), Inches(7.5)


# ---------------------------------------------------------------- pptx utils

def style_run(run, size=16, bold=False, color=INK):
    f = run.font
    f.size, f.bold, f.color.rgb, f.name = Pt(size), bold, color, FONT_LATIN
    rpr = run._r.get_or_add_rPr()
    ea = rpr.find(qn("a:ea"))
    if ea is None:
        ea = rpr.makeelement(qn("a:ea"), {})
        rpr.append(ea)
    ea.set("typeface", FONT_EA)


def text(slide, x, y, w, h, lines, size=16, color=INK, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=1.15):
    """lines: str, or list of str / (str, dict(size,bold,color,bullet))."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    if isinstance(lines, str):
        lines = [lines]
    for i, item in enumerate(lines):
        s, opt = (item, {}) if isinstance(item, str) else item
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        p.space_after = Pt(opt.get("after", 6))
        prefix = "•  " if opt.get("bullet") else ""
        r = p.add_run()
        r.text = prefix + s
        style_run(r, opt.get("size", size), opt.get("bold", False),
                  opt.get("color", color))
    return tb


def box(slide, x, y, w, h, label, fill=LIGHT, line=None, size=13,
        color=INK, bold=False, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1.25)
    s.shadow.inherit = False
    tf = s.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.06)
    for i, ln in enumerate(label.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = ln
        style_run(r, size if i == 0 else size - 2, bold and i == 0,
                  color if i == 0 else MUTED if fill == LIGHT else color)
    return s


def arrow(slide, x1, y1, x2, y2, color=MUTED, width=1.5, label=None,
          label_dx=Inches(0.08)):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    c.line.color.rgb = color
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med",
                                            "len": "med"})
    ln.append(tail)
    if label:
        text(slide, min(x1, x2) + label_dx, (y1 + y2) // 2 - Inches(0.17),
             Inches(2.6), Inches(0.4), label, size=11, color=MUTED)
    return c


def table(slide, x, y, w, rows, col_w=None, size=13, row_h=Inches(0.42),
          header_fill=INK, highlight=None):
    """rows[0] is the header. highlight: {(r, c): RGBColor} text colors."""
    highlight = highlight or {}
    n, m = len(rows), len(rows[0])
    gt = slide.shapes.add_table(n, m, x, y, w, row_h * n)
    t = gt.table
    if col_w:
        tot = sum(col_w)
        for j, cw in enumerate(col_w):
            t.columns[j].width = Emu(int(w * cw / tot))
    for i in range(n):
        t.rows[i].height = row_h
        for j in range(m):
            cell = t.cell(i, j)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = (header_fill if i == 0 else
                                        WHITE if i % 2 else LIGHT)
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
            r = p.add_run()
            r.text = str(rows[i][j])
            style_run(r, size, bold=(i == 0),
                      color=WHITE if i == 0 else highlight.get((i, j), INK))
    return gt


class Deck:

    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = W, H
        self.n = 0

    def slide(self, title, kicker=None):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.n += 1
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.12), H)
        bar.fill.solid()
        bar.fill.fore_color.rgb = ACCENT
        bar.line.fill.background()
        if kicker:
            text(s, Inches(0.55), Inches(0.28), Inches(12), Inches(0.35),
                 kicker, size=13, color=ACCENT)
        text(s, Inches(0.55), Inches(0.55), Inches(12.3), Inches(0.8), title,
             size=28, color=INK)
        self._title_run_bold(s)
        text(s, Inches(12.2), Inches(7.0), Inches(0.9), Inches(0.35),
             str(self.n), size=11, color=MUTED, align=PP_ALIGN.RIGHT)
        return s

    @staticmethod
    def _title_run_bold(s):
        tb = s.shapes[-1]
        for r in tb.text_frame.paragraphs[0].runs:
            r.font.bold = True

    def save(self, path):
        self.prs.save(path)


def note(slide, s, y=Inches(6.55), color=MUTED, size=12):
    text(slide, Inches(0.55), y, Inches(12.2), Inches(0.6), s, size=size,
         color=color)


def callout(slide, x, y, w, h, title, body, color=ACCENT):
    b = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    b.fill.solid()
    b.fill.fore_color.rgb = LIGHT
    b.line.fill.background()
    edge = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Inches(0.07), h)
    edge.fill.solid()
    edge.fill.fore_color.rgb = color
    edge.line.fill.background()
    lines = [(title, {"bold": True, "size": 15, "color": color})]
    lines += [(t, {"size": 13, "bullet": True}) for t in body]
    text(slide, x + Inches(0.2), y + Inches(0.12), w - Inches(0.3),
         h - Inches(0.2), lines)


# ---------------------------------------------------------------- data

def load_memory():
    out = {}
    for tag in ["default", "20-36", "36-20"]:
        log = os.path.join(RES, "sweep", tag, "server.log")
        if os.path.exists(log):
            parsed = parse_log(log)
            out[tag] = (parsed["num_gpu_blocks"], summarize(parsed))
    return out


def load_bench(base):
    out = {}
    for f in glob.glob(os.path.join(base, "*", "*", "bench.json")):
        part, wl = f.split(os.sep)[-3:-1]
        d = json.load(open(f))
        out[(part, wl)] = d
    return out


def load_grid_memory():
    out = {}
    for tag, _ in GRID_PARTS:
        log = os.path.join(RES2, "grid", tag, "server.log")
        if os.path.exists(log):
            parsed = parse_log(log)
            out[tag] = (parsed["num_gpu_blocks"], summarize(parsed),
                        os.path.exists(os.path.join(RES2, "grid", tag,
                                                    "FAILED")))
    return out


def load_repeats():
    """{(part, workload): [bench, ...]} over results/repeat/r*."""
    out = {}
    for r in sorted(glob.glob(os.path.join(REP, "r*"))):
        for k, v in load_bench(r).items():
            out.setdefault(k, []).append(v)
    return out


def run_analyze_profile(py, part, tmp, base=RES):
    """Per-layer cost and stage balance via hybrid_pp.analyze_profile."""
    res = {}
    for d in sorted(glob.glob(os.path.join(base, "profile", part, "*",
                                           "profile"))):
        wl = d.split(os.sep)[-2]
        if len(glob.glob(os.path.join(d, "*.jsonl"))) < 2:
            continue
        out = os.path.join(tmp, os.path.basename(base), part, wl)
        r = subprocess.run([py, "-m", "hybrid_pp.analyze_profile", d,
                            "--out", out], cwd=ROOT, capture_output=True,
                           text=True)
        if r.returncode == 0 and os.path.exists(
                os.path.join(out, "stages.csv")):
            res[wl] = (pd.read_csv(os.path.join(out, "stages.csv")),
                       pd.read_csv(os.path.join(out, "layer_cost.csv")))
    return res


# ---------------------------------------------------------------- figures

def fig_layers(path):
    parts = [("默认 28,28", 28), ("20,36", 20), ("36,20", 36)]
    colors = {"M": C_MAMBA, "-": C_MLP, "*": C_ATTN}
    fig, ax = plt.subplots(figsize=(12, 2.9), dpi=200)
    for row, (name, cut) in enumerate(parts):
        y = len(parts) - 1 - row
        for i, c in enumerate(PATTERN):
            ax.add_patch(plt.Rectangle((i, y + 0.12), 0.9, 0.76,
                                       color=colors[c], lw=0))
        ax.plot([cut - 0.05, cut - 0.05], [y - 0.02, y + 1.02], color="#111",
                lw=2.4)
        a0, a1 = PATTERN[:cut].count("*"), PATTERN[cut:].count("*")
        ax.text(-1, y + 0.5, name, ha="right", va="center", fontsize=12)
        ax.text(57, y + 0.5, f"attention：{a0} + {a1}", ha="left",
                va="center", fontsize=12,
                color="#B91C1C" if a0 != a1 else "#111")
    for i in (14, 21, 30, 39):
        ax.text(i + 0.45, 3.08, str(i), ha="center", fontsize=10,
                color=C_ATTN)
    ax.set_xlim(-8, 66)
    ax.set_ylim(-0.1, 3.35)
    ax.axis("off")
    handles = [plt.Rectangle((0, 0), 1, 1, color=c) for c in
               (C_MAMBA, C_MLP, C_ATTN)]
    ax.legend(handles, ["Mamba ×27", "MLP ×25", "Attention ×4"],
              loc="lower center", bbox_to_anchor=(0.45, -0.22), ncol=3,
              frameon=False, fontsize=12)
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)


def fig_memory(mem, path):
    order = [("default", "默认 28,28"), ("20-36", "20,36"), ("36-20", "36,20")]
    fig, ax = plt.subplots(figsize=(8.2, 4.6), dpi=200)
    xs, labels = [], []
    x = 0
    for tag, name in order:
        if tag not in mem:
            continue
        _, rows = mem[tag]
        for r in rows:
            parts = [(r["weights_gib"], C_W), (r["activation_gib"], C_ACT),
                     (r["kv_used_gib"], C_KV), (r["kv_unused_gib"], C_UNUSED)]
            bottom = 0
            for v, c in parts:
                v = max(v or 0, 0)
                ax.bar(x, v, bottom=bottom, color=c, width=0.75,
                       hatch="//" if c == C_UNUSED else None,
                       edgecolor="white", lw=0.6)
                bottom += v
            if (r["kv_unused_gib"] or 0) > 0.5:
                ax.text(x, bottom + 0.25, f"浪费\n{r['kv_unused_gib']:.1f}G",
                        ha="center", fontsize=10, color="#B91C1C")
            xs.append(x)
            labels.append(f"S{r['pp_rank']}\n({r['pattern'].count('*')}×attn)")
            x += 1
        ax.text(x - 1.5, -4.3, name, ha="center", fontsize=12)
        x += 0.6
    ax.set_xticks(xs, labels, fontsize=10)
    ax.set_ylabel("GiB（0.9 × 23.5 GiB = 21.2 GiB 可用）")
    ax.set_ylim(0, 23)
    handles = [plt.Rectangle((0, 0), 1, 1, color=c) for c in
               (C_W, C_ACT, C_KV, C_UNUSED)]
    ax.legend(handles, ["权重", "激活峰值（含 Mamba state）", "KV 实际使用",
                        "KV 预算中没用上的"], ncol=4, frameon=False,
              fontsize=9.5, loc="lower center", bbox_to_anchor=(0.5, 1.0))
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)


def fig_layer_cost(cost, path):
    pre = cost[cost.phase == "prefill"]
    fig, ax = plt.subplots(figsize=(6.2, 4.2), dpi=200)
    for kind, c, name in [("mamba", C_MAMBA, "Mamba"), ("mlp", C_MLP, "MLP"),
                          ("attention", C_ATTN, "Attention")]:
        d = pre[pre.kind == kind].sort_values("bucket")
        ax.plot(d.bucket, d.mean_ms, "o-", color=c, label=name, lw=2)
    ax.set_xscale("log", base=2)
    ax.set_yscale("log")
    ax.set_xlabel("这一步的 prefill token 数")
    ax.set_ylabel("每层 GPU 耗时（ms）")
    ax.legend(frameon=False)
    ax.grid(alpha=0.3, which="both")
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)


def fig_grid(grid, path):
    """Request throughput of each split relative to 28,28, per workload."""
    fig, ax = plt.subplots(figsize=(8.0, 4.6), dpi=200)
    xs = [n for _, n in GRID_PARTS]
    for (wl, name), c, mk in zip(GRID_WLS, C_SERIES, "osD^"):
        base = grid[("default", wl)]["request_throughput"]
        pts = [(n, grid[(t, wl)]["request_throughput"] / base * 100)
               for t, n in GRID_PARTS if (t, wl) in grid]
        ax.plot(*zip(*pts), marker=mk, color=c, lw=2, ms=8, label=name)
        ax.text(pts[-1][0] + 0.5, pts[-1][1], name.split(" ")[0],
                color="#374151", fontsize=10, va="center")
    for n in (16, 40):
        ax.text(n, 96.5, "起不来\n(KV 不足)", ha="center", fontsize=10,
                color="#B91C1C")
    ax.axhline(100, color="#9CA3AF", lw=1, ls="--")
    ax.set_xticks(xs, [f"{n},{56 - n}\n({PATTERN[:n].count('*')}+"
                       f"{PATTERN[n:].count('*')})" for n in xs], fontsize=10)
    ax.set_xlabel("切分（括号内为 attention 层数 S0+S1）")
    ax.set_ylabel("吞吐，相对默认 28,28（%）")
    ax.set_xlim(14, 43.5)
    ax.set_ylim(74, 118)
    ax.grid(axis="y", alpha=0.3)
    ax.legend(frameon=False, fontsize=10, ncol=4, loc="lower center",
              bbox_to_anchor=(0.5, 1.0))
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)


def fig_capacity(gmem, path):
    """KV token capacity per split, with the two thresholds that matter."""
    fig, ax = plt.subplots(figsize=(7.4, 4.4), dpi=200)
    xs, labels = [], []
    for i, (tag, n) in enumerate(GRID_PARTS):
        if tag not in gmem:
            continue
        blocks, rows, failed = gmem[tag]
        cap = blocks * 16 / 1000
        ax.bar(i, cap, width=0.7, color=C_KV if not failed else "#F87171")
        waste = max(max(r["kv_unused_gib"] or 0, 0) for r in rows)
        ax.text(i, cap + 8, f"{cap:.0f}K\n浪费 {waste:.1f}G", ha="center",
                fontsize=9.5, color="#374151")
        xs.append(i)
        labels.append(f"{n},{56 - n}")
    ax.axhline(264, color="#EA580C", lw=1.5, ls="--",
               label="16K 负载需要 ≈264K（16 并发 × 16.5K）")
    ax.axhline(65.5, color="#B91C1C", lw=1.2, ls=":",
               label="启动下限 65.5K（每个 VE 放下一条 32K 序列）")
    ax.legend(frameon=False, fontsize=9.5, loc="lower center",
              bbox_to_anchor=(0.5, 1.0), ncol=1)
    ax.set_xticks(xs, labels)
    ax.set_ylabel("KV 可容纳 token 数（千）")
    ax.set_ylim(0, 590)
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)


# ---------------------------------------------------------------- slides

def build(args):
    tmp = tempfile.mkdtemp()
    mem = load_memory()
    bench = load_bench(os.path.join(RES, "sweep"))
    prof = {p: run_analyze_profile(args.profile_python, p, tmp)
            for p in ["default", "36-20"]}
    # day 1 lost default 16K to an OOM; day 2 reran it at util 0.85
    prof["default"].update(run_analyze_profile(args.profile_python, "default",
                                               tmp, base=RES2))
    prof_bench = load_bench(os.path.join(RES, "profile"))
    prof_bench.update(load_bench(os.path.join(RES2, "profile")))
    gmem = load_grid_memory()
    grid = load_bench(os.path.join(RES2, "grid"))
    reps = load_repeats()

    f_layers = os.path.join(tmp, "layers.png")
    f_mem = os.path.join(tmp, "mem.png")
    f_cost = os.path.join(tmp, "cost.png")
    f_grid = os.path.join(tmp, "grid.png")
    f_cap = os.path.join(tmp, "cap.png")
    fig_grid(grid, f_grid)
    fig_capacity(gmem, f_cap)
    fig_layers(f_layers)
    fig_memory(mem, f_mem)
    cost_all = pd.concat([c for _, c in prof["default"].values()]) \
        if prof["default"] else None
    if cost_all is not None:
        cost_all = (cost_all.groupby(["phase", "kind", "bucket"])
                    .apply(lambda g: pd.Series({
                        "mean_ms": (g.mean_ms * g.samples).sum()
                        / g.samples.sum(),
                        "samples": g.samples.sum()}))
                    .reset_index())
        fig_layer_cost(cost_all, f_cost)

    d = Deck()

    # 1 title
    s = d.prs.slides.add_slide(d.prs.slide_layouts[6])
    d.n += 1
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = INK
    bg.line.fill.background()
    text(s, Inches(0.9), Inches(2.2), Inches(11.5), Inches(1.2),
         "混合模型流水线并行推理：阶段性验证结果", size=36, color=WHITE)
    s.shapes[-1].text_frame.paragraphs[0].runs[0].font.bold = True
    text(s, Inches(0.9), Inches(3.35), Inches(11.5), Inches(1.4),
         ["Nemotron-Nano-9B-v2（Mamba-Transformer 混合）× vLLM 0.10.1.1 × 2×RTX 4090",
          ("vLLM 的 PP bug 复现与修复 · 正确性验证 · 显存与计算不均衡 · 7 种切分 × 4 种负载",
           {"color": RGBColor(0xC7, 0xD2, 0xFE)})],
         size=18, color=RGBColor(0xE5, 0xE7, 0xEB))
    text(s, Inches(0.9), Inches(6.3), Inches(11), Inches(0.5),
         "2026-10-28 · 进度汇报", size=14, color=RGBColor(0x9C, 0xA3, 0xAF))

    # 2 summary
    s = d.slide("要回答的 6 个问题", "一页结论")
    rows = [["#", "问题", "结果", "对论文的意义"],
            ["0", "旧驱动 / 4090 能不能用", "✅ 可用（需 NCCL_P2P_DISABLE=1）",
             "不用换机器"],
            ["1", "vLLM 的 Nemotron PP 真的有 bug 吗",
             "✅ 4 个 bug，全部真机复现", "现有框架对混合模型 PP 支持不成熟"],
            ["2", "修复后结果正确吗", "✅ 10/10 逐字一致（2 种切分）",
             "后续所有实验的前提"],
            ["3", "显存浪费存在吗", "✅ 不均匀切分浪费 5.8 GiB（预算的 90%）",
             "图 2：显存利用率低"],
            ["4", "计算不均衡存在吗", "⚠️ 存在，但主要来自层数而非 attention",
             "图 1/3 需要更长上下文"],
            ["5", "最优静态切分会随负载变吗", "✅ 会：短输入偏 36,20（+8～14%），"
             "长输入偏 28,28（36,20 掉 ~23%）", "图 4：静态切分不够用"]]
    table(s, Inches(0.55), Inches(1.55), Inches(12.2), rows,
          col_w=[0.5, 3.2, 4.3, 4.2], size=15, row_h=Inches(0.62),
          highlight={(5, 2): WARN, (6, 2): GOOD})
    note(s, "PP=4（无 attention 的 stage）需要 4 卡，尚未做。"
            "关键吞吐点已重复 3 次，波动 < 2%。", y=Inches(6.65))

    # 3 setup
    s = d.slide("实验平台与软件栈", "环境")
    rows = [["项目", "配置"],
            ["GPU", "2 × NVIDIA RTX 4090 24 GB（PCIe，无 NVLink；NCCL_P2P_DISABLE=1）"],
            ["推理框架", "vLLM 0.10.1.1，V0 引擎（VLLM_USE_V1=0），流水线并行 PP=2"],
            ["模型", "nvidia/NVIDIA-Nemotron-Nano-9B-v2，bf16，56 层"],
            ["关键参数", "max-model-len 32768 · gpu-memory-utilization 0.9 · "
                      "max-num-seqs 16 · Mamba SSM cache fp32 · block 16 token"],
            ["我们的改动", "不改 vLLM 源码：一个 vLLM 插件（vllm.general_plugins）"
                       "在每个进程里打补丁"],
            ["插件内容", "4 个 bug 修复（可单独开关）· 每 stage 显存明细日志 · "
                      "CUDA event 逐层计时 profiler"],
            ["负载", "vllm bench serve，随机数据集，输入 1K～30K token，固定并发"]]
    table(s, Inches(0.55), Inches(1.55), Inches(12.2), rows, col_w=[2, 10],
          size=15, row_h=Inches(0.56))
    note(s, "为什么用 V0：V0 的 PP 实现是 virtual engine + 每卡独立算 KV block 再取 min()，"
            "正是我们要研究的“静态切分”行为；V1 对混合模型的 PP 支持更不完整。")

    # 4 model structure
    s = d.slide("模型结构：4 层 attention 稀疏地分布在 56 层里", "Nemotron-Nano-9B-v2")
    s.shapes.add_picture(f_layers, Inches(0.55), Inches(1.5),
                         width=Inches(12.2))
    callout(s, Inches(0.55), Inches(5.0), Inches(5.9), Inches(1.85),
            "三类层的成本特点",
            ["Mamba / MLP：耗时只看本步 token 数，与上下文长度无关",
             "Attention：耗时随上下文增长；KV 4 KiB/token/层",
             "Mamba state：每层每序列 5.1 MiB，按 max_num_seqs 预分配"])
    callout(s, Inches(6.85), Inches(5.0), Inches(5.9), Inches(1.85),
            "三种切分的设计（控制变量）",
            ["默认 28,28：attention 2+2，作对照组",
             "20,36 和 36,20：1+3 与 3+1，互为镜像",
             "若现象随 attention 一起换边 → 原因是 attention 分布"],
            color=WARN)

    # 5 architecture
    s = d.slide("vLLM V0 的 PP 调用链", "架构（详见 docs/vllm_callchain.md）")
    X = Inches
    box(s, X(0.6), X(1.6), X(2.2), X(0.8), "HTTP 请求\n/v1/completions",
        fill=LIGHT)
    box(s, X(3.4), X(1.6), X(3.0), X(0.8),
        "AsyncLLMEngine\nrun_engine_loop", fill=ACCENT, color=WHITE, bold=True)
    arrow(s, X(2.8), X(2.0), X(3.4), X(2.0))
    box(s, X(7.2), X(1.25), X(2.6), X(0.7), "Scheduler VE0\nKV: N/2 block",
        fill=LIGHT, line=ACCENT)
    box(s, X(7.2), X(2.15), X(2.6), X(0.7), "Scheduler VE1\nKV: N/2 block",
        fill=LIGHT, line=ACCENT)
    arrow(s, X(6.4), X(1.85), X(7.2), X(1.6))
    arrow(s, X(6.4), X(2.15), X(7.2), X(2.5))
    text(s, X(10.0), X(1.3), X(3.0), X(1.5),
         ["PP=2 → 2 个 virtual engine",
          "新请求进未完成数最少的 VE",
          "每个 VE 同时只有 1 个 micro-batch"], size=12, color=MUTED)
    box(s, X(3.4), X(3.2), X(3.0), X(0.8),
        "Executor（多进程）\n每个 stage 一把 asyncio.Lock", fill=INK, color=WHITE,
        bold=True)
    arrow(s, X(4.9), X(2.4), X(4.9), X(3.2), label="micro-batch")
    box(s, X(0.6), X(4.6), X(4.6), X(1.35),
        "Stage 0 · GPU0 · 层 0–27\nembedding → 本 stage 各层\n"
        "Mamba cache + KV（VE0/VE1 各一份）", fill=RGBColor(0xCC, 0xFB, 0xF1),
        line=RGBColor(0x0F, 0x76, 0x6E), bold=True)
    box(s, X(7.2), X(4.6), X(4.6), X(1.35),
        "Stage 1 · GPU1 · 层 28–55\n本 stage 各层 → norm → logits → 采样\n"
        "Mamba cache + KV（VE0/VE1 各一份）", fill=RGBColor(0xCC, 0xFB, 0xF1),
        line=RGBColor(0x0F, 0x76, 0x6E), bold=True)
    arrow(s, X(4.2), X(4.0), X(2.9), X(4.6))
    arrow(s, X(5.6), X(4.0), X(9.5), X(4.6))
    arrow(s, X(5.2), X(5.27), X(7.2), X(5.27), color=ACCENT, width=2.5)
    text(s, X(5.25), X(4.75), X(1.95), X(0.5),
         "hidden_states\n+ residual", size=11, color=ACCENT,
         align=PP_ALIGN.CENTER)
    arrow(s, X(11.8), X(5.0), X(11.8), X(2.4), color=MUTED)
    text(s, X(11.9), X(3.3), X(1.3), X(0.6), "采样结果\n回到引擎", size=11,
         color=MUTED)
    note(s, "启动时：每个 worker 用假数据跑一次 forward 测显存 → 各自算出能放多少 KV block → "
            "executor 取 min() → 每个 VE 再分到 1/PP。", y=Inches(6.3),
         color=INK, size=13)

    # 6 bugs
    s = d.slide("发现 ①：vLLM 的 NemotronH PP 有 4 个 bug，一个挡一个", "可直接写进 motivation")
    rows = [["Bug", "问题", "4090 上的现象", "状态"],
            ["3", "load_weights 不跳过其他 stage 的层",
             "加载时 KeyError: 'layers.49.mixer…'", "✅ 复现 + 修复"],
            ["1", "forward 遍历全部 56 层（含占位符 PPMissingLayer）",
             "ValueError: too many values to unpack", "✅ 复现 + 修复"],
            ["4", "工厂函数名写错（intmd vs intermediate），\n接口里的空方法静默返回 None",
             "stage 1 上 AssertionError", "✅ 复现 + 修复"],
            ["2", "非首 stage 把收到的 residual 清空",
             "能启动，但 10/10 输出错误", "✅ 复现 + 修复"]]
    table(s, Inches(0.55), Inches(1.55), Inches(12.2), rows,
          col_w=[0.7, 5.0, 4.3, 2.0], size=14, row_h=Inches(0.72))
    callout(s, Inches(0.55), Inches(5.3), Inches(12.2), Inches(1.4),
            "怎么证实的：每次多修一个，逐层剥开",
            ["原版 → 挂在加载（Bug 3）；修 3 → stage 0 挂在 Bug 1，同时 stage 1 挂在 Bug 4"
             "（两张卡并行 profile，日志里两个错误都出现了）；修 3+1+4 → 能跑但结果错（Bug 2）；全修 → 正确",
             "v0.10.2 / v0.11.0 中这些代码仍然存在；全部通过插件修复，不改 vLLM 源码"])

    # 7 correctness
    s = d.slide("发现 ②：修复后 PP=2 与单卡逐字一致", "后续所有实验的前提 ✅")
    rows = [["配置", "与 PP=1 一致", "第一个不同的字符位置"],
            ["PP=2 默认 28,28（全部修复）", "10 / 10", "—"],
            ["PP=2 不均匀 20,36（全部修复）", "10 / 10", "—"],
            ["PP=2 保留 Bug 2（残差被清空）", "0 / 10", "第 0～5 个字符"]]
    table(s, Inches(0.55), Inches(1.55), Inches(7.4), rows,
          col_w=[3.6, 1.6, 2.2], size=14, row_h=Inches(0.55),
          highlight={(1, 1): GOOD, (2, 1): GOOD, (3, 1): BAD})
    callout(s, Inches(8.3), Inches(1.55), Inches(4.45), Inches(2.2),
            "测试方法",
            ["temperature=0（greedy），输出确定",
             "10 条 prompt：问答 / 代码 / 中译英 / 算术 / ~3000 token 长文",
             "长 prompt 覆盖 Mamba 多 chunk prefill 与跨多个 KV block"])
    text(s, Inches(0.55), Inches(4.05), Inches(12.2), Inches(0.4),
         "Bug 2 的输出示例（prompt：“def fibonacci(n):”）", size=14,
         color=MUTED)
    text(s, Inches(0.55), Inches(4.45), Inches(12.2), Inches(1.4),
         [("PP=1：  if n <= 0: return 0 elif n == 1: return 1 else: return fibonacci(n-1)…",
           {"color": GOOD}),
          ("Bug 2：the the the\\nonly\\n\\n to only\\n00 to the …",
           {"color": BAD})], size=15)
    note(s, "意义：可以任意改变切分而不影响正确性 —— 后面扫描切分的实验靠的就是这一点。")

    # 8 memory mechanism
    s = d.slide("显存是怎么分配的：每卡各算一个数，然后取最小值", "机制（vLLM 源码）")
    steps = [("① 每个 stage 测显存", "用假数据跑一次 forward\n得到 权重 + 激活峰值"),
             ("② 算本卡能放多少 block", "(0.9×总显存 − 非KV) ÷ block 大小\nblock 大小 ∝ 本卡 attention 层数"),
             ("③ executor 取 min()", "所有 stage 都只分配\n最小的那个 block 数"),
             ("④ 再按 VE 均分", "每个 virtual engine\n只拿 1/PP")]
    for i, (a, b) in enumerate(steps):
        x = Inches(0.55 + i * 3.1)
        box(s, x, Inches(1.6), Inches(2.8), Inches(1.5), a + "\n" + b,
            fill=LIGHT, line=ACCENT, size=15, bold=True)
        if i < 3:
            arrow(s, x + Inches(2.8), Inches(2.35), x + Inches(3.1),
                  Inches(2.35), color=ACCENT)
    callout(s, Inches(0.55), Inches(3.5), Inches(6.0), Inches(2.6),
            "为什么混合模型会浪费",
            ["attention 只有 1 层的 stage：block 只有 64 KiB，本来能放 10 万个",
             "attention 3 层且权重更多的 stage：block 192 KiB，只能放 1 万个",
             "取 min() 后两边都只用 1 万个 → 前者 90% 的 KV 预算闲置"],
            color=BAD)
    callout(s, Inches(6.75), Inches(3.5), Inches(6.0), Inches(2.6),
            "容量由什么决定（不是“浪费”本身）",
            ["可容纳 token 数 = min( KV 预算ᵢ ÷ 每 token 的 KVᵢ )",
             "空出来的显存放不了另一张卡那几层的 KV",
             "→ 切分同时决定了每张卡的显存供给和 KV 需求，按层数均分会让两者错位"])

    # 9 memory results
    s = d.slide("发现 ③：不均匀切分时 90% 的 KV 预算闲置，且随切分换边", "图 2 雏形")
    s.shapes.add_picture(f_mem, Inches(0.45), Inches(1.45),
                         height=Inches(5.0))
    mrows = [["切分", "S0 / S1 浪费", "engine block", "可容纳 token"]]
    for tag, name in [("default", "28,28"), ("20-36", "20,36"),
                      ("36-20", "36,20")]:
        if tag in mem:
            blocks, rows_ = mem[tag]
            mrows.append([name, " / ".join(f"{max(r['kv_unused_gib'], 0):.2f}"
                                           for r in rows_),
                          f"{blocks:,}", f"{blocks * 16 / 1000:.0f}K"])
    table(s, Inches(9.25), Inches(1.6), Inches(3.75), mrows,
          col_w=[1.0, 1.5, 1.3, 1.2], size=12, row_h=Inches(0.48),
          highlight={(2, 1): BAD, (3, 1): BAD})
    text(s, Inches(9.25), Inches(3.75), Inches(3.75), Inches(2.6),
         [("对照组几乎不浪费（0.18 GiB）", {"bullet": True}),
          ("1+3 / 3+1：5.8 GiB 闲置，位置互换", {"bullet": True}),
          ("总容量 516K → 161K token（少 3.2 倍）", {"bullet": True, "bold": True}),
          ("激活峰值 7.2～7.8 GiB：其中 32K token 的 profile prefill 占大头",
           {"bullet": True})], size=13)
    note(s, "max-num-seqs=64 时 20,36 和 36,20 直接 OOM 起不来（Mamba state 预分配 5.4 GiB）；"
            "上表为 max-num-seqs=16。", y=Inches(6.65))

    # 9b memory grid
    s = d.slide("显存：7 种切分扫描 —— 容量差 10 倍，两端直接起不来",
                "图 2 · results/day2/grid")
    s.shapes.add_picture(f_cap, Inches(0.45), Inches(1.45), height=Inches(4.9))
    grows = [["切分", "attn", "S0 / S1 浪费 GiB", "容量"]]
    for tag, n in GRID_PARTS:
        if tag not in gmem:
            continue
        blocks, rows_, failed = gmem[tag]
        grows.append([f"{n},{56 - n}",
                      f"{PATTERN[:n].count('*')}+{PATTERN[n:].count('*')}",
                      " / ".join(f"{max(r['kv_unused_gib'], 0):.2f}"
                                 for r in rows_),
                      "起不来" if failed else f"{blocks * 16 / 1000:.0f}K"])
    table(s, Inches(8.6), Inches(1.5), Inches(4.4), grows,
          col_w=[0.9, 0.7, 1.75, 1.05], size=12, row_h=Inches(0.4),
          highlight={(1, 3): BAD, (7, 3): BAD, (4, 3): GOOD})
    text(s, Inches(8.6), Inches(4.85), Inches(4.4), Inches(1.7),
         [("只有 28,28 两边“供需”刚好对上", {"bullet": True}),
          ("40,16：4 层 attention 全在 S0，S1 的 7.2 GiB 整块闲置",
           {"bullet": True}),
          ("16,40 / 40,16：权重多的一边只剩 0.6～0.8 GiB 给 KV，"
           "连一条 32K 序列都放不下", {"bullet": True, "bold": True})],
         size=12)

    # 10 throughput: repeats
    s = d.slide("吞吐：重复 3 次 —— 短输入 36,20 更快是真的",
                "results/repeat · 每组 64 条请求 × 3 次")

    def ms(vals, f="{:.0f}"):
        m = sum(vals) / len(vals)
        sd = (sum((v - m) ** 2 for v in vals) / max(len(vals) - 1, 1)) ** 0.5
        return m, (f + " ± " + f).format(m, sd)

    trows = [["切分", "1K：输出 tok/s", "1K TPOT (ms)", "16K：输出 tok/s",
              "16K TTFT (s)"]]
    tp = {}
    for tag, name in [("default", "默认 28,28"), ("36-20", "36,20")]:
        a, b = reps.get((tag, "1024-256"), []), reps.get((tag, "16384-128"), [])
        if not (a and b):
            continue
        tp[tag] = (ms([x["output_throughput"] for x in a])[0],
                   ms([x["output_throughput"] for x in b])[0])
        trows.append([name, ms([x["output_throughput"] for x in a])[1],
                      ms([x["mean_tpot_ms"] for x in a], "{:.1f}")[1],
                      ms([x["output_throughput"] for x in b], "{:.1f}")[1],
                      ms([x["mean_ttft_ms"] / 1000 for x in b], "{:.1f}")[1]])
    table(s, Inches(0.55), Inches(1.55), Inches(12.2), trows,
          col_w=[2.4, 2.6, 2.2, 2.6, 2.4], size=15, row_h=Inches(0.55),
          highlight={(2, 1): GOOD, (1, 3): GOOD, (2, 4): BAD})
    d1 = tp["36-20"][0] / tp["default"][0] - 1 if len(tp) == 2 else 0
    d16 = tp["36-20"][1] / tp["default"][1] - 1 if len(tp) == 2 else 0
    callout(s, Inches(0.55), Inches(3.5), Inches(6.0), Inches(2.9),
            f"短输入：36,20 快 {d1:.0%}，标准差 < 1% → 不是噪声",
            ["但 eager profiling 里 S0 单步耗时是 S1 的 2.3 倍，按理应更慢",
             "关键：profiler 只计 model.forward，不含 lm_head 和采样",
             "这两步只在最后一个 stage 做：lm_head 是 4480×131072 的矩阵，"
             "decode 时约等于 4 层的计算量，再加上采样",
             "→ 末 stage 天然更重，“均匀”的 28,28 其实不均匀（待测量验证）"],
            color=GOOD)
    callout(s, Inches(6.75), Inches(3.5), Inches(6.0), Inches(2.9),
            f"长输入：36,20 慢 {-d16:.0%}，TTFT 从 12.8 s 涨到 21.9 s",
            ["16 并发 × 16.5K ≈ 264K token，36,20 只能放 161K",
             "请求排队等 KV → 显存不均衡直接变成吞吐损失",
             "同一个切分在两种负载下一好一坏 → 下一页的完整扫描"],
            color=BAD)
    note(s, "eager 模式（关 CUDA graph）下 36,20 反而慢 36%：多 8 层带来的 kernel 启动开销"
            "在 CPU 上串行，CUDA graph 下这部分被消掉了（推测）。", y=Inches(6.6))

    # 10b grid
    s = d.slide("发现 ⑤：最优的静态切分随负载变化", "图 4 · 7 种切分 × 4 种负载（day2）")
    s.shapes.add_picture(f_grid, Inches(0.45), Inches(1.4), height=Inches(5.0))
    best = []
    for wl, name in GRID_WLS:
        cand = [(grid[(t, wl)]["request_throughput"], n) for t, n in GRID_PARTS
                if (t, wl) in grid]
        top, worst = max(cand), min(cand)
        base = grid[("default", wl)]["request_throughput"]
        best.append([name.split(" ")[0], f"{top[1]},{56 - top[1]}",
                     f"{top[0] / base - 1:+.0%}",
                     f"{worst[1]},{56 - worst[1]}  {worst[0] / base - 1:+.0%}"])
    table(s, Inches(8.45), Inches(1.5), Inches(4.5),
          [["负载", "最优", "比默认", "最差"]] + best,
          col_w=[1.1, 0.9, 0.9, 1.6], size=12, row_h=Inches(0.45))
    # Best single split across all workloads: the one whose worst share of
    # each workload's own optimum is highest. Bounds what dynamic can gain.
    def share(t, wl):
        top = max(grid[(u, wl)]["request_throughput"] for u, _ in GRID_PARTS
                  if (u, wl) in grid)
        return grid[(t, wl)]["request_throughput"] / top
    comp = max((min(share(t, wl) for wl, _ in GRID_WLS), n)
               for t, n in GRID_PARTS
               if all((t, wl) in grid for wl, _ in GRID_WLS))
    text(s, Inches(8.45), Inches(3.95), Inches(4.5), Inches(2.6),
         [("decode 为主（短输入、长输出）：S0 层数越多越快，"
           "36,20 比默认快 10～14%", {"bullet": True}),
          ("prefill / 长上下文为主：28,28 最好，偏离越远越差（KV 容量）",
           {"bullet": True}),
          ("没有一个静态切分在所有负载下都最优 → 支撑“按负载动态调整切分”",
           {"bullet": True, "bold": True, "color": ACCENT}),
          (f"但折中切分 {comp[1]},{56 - comp[1]} 在 4 种负载下都不低于最优的 "
           f"{comp[0]:.0%}（最多差 {1 - comp[0]:.0%}）：动态切分的收益空间"
           "需要在更长上下文和 PP=4 上确认",
           {"bullet": True, "color": BAD})], size=12)
    note(s, "单次运行；纵轴为请求吞吐（输入输出长度固定，等价于 token 吞吐）。"
            "16,40 与 40,16 启动即因 KV 放不下 32K 序列失败。", y=Inches(6.6))

    # 11 profiling
    s = d.slide("发现 ④：计算不均衡 —— 目前主要来自层数，而非 attention", "图 1 / 图 3 雏形 · 部分结果")
    if cost_all is not None:
        s.shapes.add_picture(f_cost, Inches(0.45), Inches(1.45),
                             height=Inches(4.2))
        pre = cost_all[cost_all.phase == "prefill"].pivot(
            index="bucket", columns="kind", values="mean_ms")
        share = []
        for b in [1024, 8192, 16384]:
            if b in pre.index:
                r = pre.loc[b]
                tot = r.attention * 4 + r.mamba * 27 + r.mlp * 25
                share.append(f"{b // 1024}K：{r.attention * 4 / tot:.1%}")
        text(s, Inches(0.55), Inches(5.75), Inches(6.1), Inches(0.8),
             [("prefill 时 attention 占全模型耗时：" + " → ".join(share),
               {"bold": True}),
              ("单层 attention 增长最快（≈二次方），但全模型只有 4 层", {})],
             size=13)
    brows = [["切分 / 输入", "S0 单步 ms", "S1 单步 ms", "S0 忙碌率", "S1 忙碌率"]]
    ratios = []
    for part, pname in [("default", "28,28"), ("36-20", "36,20")]:
        for wl in ["1024-128", "8192-128", "16384-128", "30720-64"]:
            if wl not in prof[part]:
                continue
            st = prof[part][wl][0].set_index("pp_rank")
            b = prof_bench.get((part, wl))
            if b and b.get("completed", 0) < 32:
                continue  # crashed run, see notes
            brows.append([f"{pname} · {int(wl.split('-')[0]) // 1024}K",
                          f"{st.loc[0, 'mean_step_ms']:.1f}",
                          f"{st.loc[1, 'mean_step_ms']:.1f}",
                          f"{st.loc[0, 'busy_ratio']:.2f}",
                          f"{st.loc[1, 'busy_ratio']:.2f}"])
            if part == "36-20":
                ratios.append(f"{st.loc[0, 'mean_step_ms'] / st.loc[1, 'mean_step_ms']:.1f}")
    table(s, Inches(6.9), Inches(1.55), Inches(5.9), brows,
          col_w=[1.7, 1.1, 1.1, 1.0, 1.0], size=12, row_h=Inches(0.4))
    text(s, Inches(6.9), Inches(1.6) + Inches(0.4) * len(brows) + Inches(0.15),
         Inches(5.9), Inches(1.8),
         [("默认切分：两卡耗时接近，但忙碌率只有 65～75%（流水线空泡）",
           {"bullet": True}),
          ("36,20：stage 0 忙 85～88%，stage 1 只有 37～44%；S0/S1 单步耗时比 "
           + " → ".join(ratios) + "，不随上下文拉大 → 差距来自 36 vs 20 层",
           {"bullet": True}),
          ("默认 16K：day1 运行中 OOM，day2 用 util 0.85 重跑补齐",
           {"bullet": True}),
          ("注意：计时不含 lm_head 和采样（只在 S1 上做），S1 的负担被低估",
           {"bullet": True, "color": BAD})], size=12)

    # 12 code reading
    s = d.slide("读源码的其他发现（调用链笔记 2.1–2.6）", "影响实验设计")
    items = [
        ("KV 再按 VE 均分", "scheduler.py:442 num_gpu_blocks //= PP：每个 virtual engine 只能用 min() 结果的一半"),
        ("Mamba state 按上限预分配", "mamba_cache.py:37 按 max_num_seqs 分配；首次 forward 时惰性创建，所以被算进“激活”"),
        ("Mamba 槽位可能不够（推测）", "每个 VE 各允许 max_num_seqs 个并发，但 Mamba 槽位共 max_num_seqs 个；并发 > 上限时可能出错，待验证"),
        ("embedding / lm_head 每卡一份", "每个 stage 多占约 1.1 GiB 用不上的权重；与异构无关，但做 baseline 时要先修掉或统一"),
        ("没有 attention 的 stage 让引擎起不来", "block 大小为 0 → 本卡 0 个 block → min()=0；PP=4 均匀切分时首尾两卡就是这样"),
    ]
    for i, (a, b) in enumerate(items):
        y = Inches(1.55 + i * 0.98)
        box(s, Inches(0.55), y, Inches(3.4), Inches(0.82), a, fill=INK,
            color=WHITE, size=14, bold=True)
        text(s, Inches(4.15), y, Inches(8.6), Inches(0.82), b, size=14,
             anchor=MSO_ANCHOR.MIDDLE)

    # 13 issues
    s = d.slide("问题与局限（如实说明）", "需要老师知道的")
    callout(s, Inches(0.55), Inches(1.55), Inches(6.0), Inches(2.4),
            "数据可信度",
            ["关键点已重复 3 次（波动 < 2%）；7×4 网格目前只跑 1 次",
             "profiling 关掉 CUDA graph，且不含 lm_head/采样 → 低估末 stage",
             "短输入 36,20 更快的原因目前是推测，需单独计时 lm_head + 采样"],
            color=WARN)
    callout(s, Inches(6.75), Inches(1.55), Inches(6.0), Inches(2.4),
            "运行中的 OOM（已绕过，但值得记录）",
            ["默认切分 16K profiling 时，stage 1 在 Mamba 长 prefill 处 OOM",
             "vLLM 启动时的显存估计偏低（碎片 2.3 GiB 未计入）",
             "已用 util 0.85 重跑补齐；说明 profile_run 对 Mamba 的估计不可靠"],
            color=BAD)
    callout(s, Inches(0.55), Inches(4.15), Inches(6.0), Inches(2.3),
            "平台带来的混杂因素",
            ["4090 无 NVLink 且关闭 P2P：stage 间传输经过主机内存",
             "16K prefill 一次要传 ~290 MB 的 hidden+residual",
             "正式实验需单独测量传输开销，或换 A100/H100"], color=MUTED)
    callout(s, Inches(6.75), Inches(4.15), Inches(6.0), Inches(2.3),
            "对论点的修正",
            ["Nano 9B v2 只有 4 层 attention，≤30K 时 attention 占耗时 <10%",
             "计算不均衡的论据需要更长上下文（64K～128K）或 attention 更多的模型",
             "显存不均衡的论据已经很强，可作为主线"], color=ACCENT)

    # 14 next
    s = d.slide("下一步计划", "接下来 1～2 周")
    rows = [["优先级", "任务", "目的"],
            ["P0", "单独计时 lm_head + 采样 + 卡间传输（CUDA graph 下）", "解释短输入 36,20 更快，修正“均衡”的定义"],
            ["P0", "网格重复 3 次；加 PP=1 单卡 ×2（数据并行）对照", "图 4 加误差；说明 PP 在什么场景才划算"],
            ["P1", "扫描 max-num-seqs × 上下文长度 × 切分", "显存不均衡的两个来源（Mamba state vs KV）"],
            ["P1", "64K～128K 长上下文", "让 attention 的计算不均衡显现出来"],
            ["P1", "测量 stage 间传输开销（P2P 关闭）", "排除平台混杂因素"],
            ["P2", "租 4 卡验证 PP=4（无 attention 的 stage）", "传统 PP 不适配混合模型的又一证据"],
            ["P2", "设计按负载调整切分的方案，与最优静态切分对比", "论文方法部分"]]
    table(s, Inches(0.55), Inches(1.55), Inches(12.2), rows,
          col_w=[1.0, 6.2, 5.0], size=14, row_h=Inches(0.58),
          highlight={(1, 0): BAD, (2, 0): BAD})

    out = args.out or os.path.join(ROOT, "report", "day1_report.pptx")
    d.save(out)
    print("wrote", out)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--profile-python", default="/venv/vllm010/bin/python",
                    help="python with vLLM deps for hybrid_pp.analyze_profile")
    ap.add_argument("--out", help="output .pptx (default report/day1_report.pptx)")
    build(ap.parse_args())
